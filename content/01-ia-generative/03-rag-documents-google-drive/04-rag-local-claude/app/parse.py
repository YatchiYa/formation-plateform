"""Du fichier au texte, puis en morceaux. Formats : PDF, Word, PowerPoint, texte, Markdown, HTML, CSV, JSON."""

import html
import io
import os
import re

INDEXABLE = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown", ".html", ".htm", ".csv", ".json"}


def extract(name: str, data: bytes) -> str:
    ext = os.path.splitext(name)[1].lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        pages = PdfReader(io.BytesIO(data)).pages
        return "\n\n".join(f"[page {i + 1}]\n{p.extract_text() or ''}" for i, p in enumerate(pages))
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:  # tableaux : une ligne par rangée, cellules séparées par « | »
            parts += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation
        slides = []
        for i, s in enumerate(Presentation(io.BytesIO(data)).slides):
            texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
            slides.append(f"[diapositive {i + 1}]\n" + "\n".join(texts))
        return "\n\n".join(slides)
    text = data.decode("utf-8", errors="replace")
    if ext in (".html", ".htm"):
        text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
        text = html.unescape(re.sub(r"<[^>]+>", " ", text))
    return text


def chunks(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    """Morceaux d'environ `size` caractères, coupés aux paragraphes puis aux phrases, avec recouvrement."""
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        return []
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    pieces = []
    for p in paras:  # un paragraphe trop long est coupé aux fins de phrases
        if len(p) <= size:
            pieces.append(p)
        else:
            pieces += [s for s in re.split(r"(?<=[.!?])\s+", p) if s]
    out, cur = [], ""
    for piece in pieces:
        while len(piece) > size:  # phrase géante (tableau aplati…) : coupe franche
            out.append(piece[:size])
            piece = piece[size - overlap:]
        if cur and len(cur) + len(piece) + 1 > size:
            out.append(cur)
            cur = cur[-overlap:] + " " + piece  # recouvrement : la fin du morceau précédent
        else:
            cur = f"{cur}\n{piece}" if cur else piece
    if cur:
        out.append(cur)
    return out
